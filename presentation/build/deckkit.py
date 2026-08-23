"""Reusable helper layer for the Synora deck (python-pptx)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml
from PIL import Image
import os

EMU_IN = 914400

# ---- Palette (Synora brand) -------------------------------------------------
INK      = "22252B"   # charcoal — dark bands / primary text on light
INK_PANEL= "2B2F37"   # lifted panel on dark
INK_SOFT = "31363F"
SAGE     = "53655C"   # brand green
SAGE_DK  = "455A50"
SAGE_SOFT= "E7ECE7"   # pale sage tint (light bg chips)
SAGE_LT  = "9FB6A6"   # light sage (on dark)
SLATE    = "A3AEB1"   # muted secondary
MIST     = "E5E8EB"   # borders/dividers
TAUPE    = "B09F95"   # warm accent
TAUPE_LT = "C9B9AE"
BG       = "FAF9F7"   # warm off-white page
SURFACE  = "FFFFFF"
TEXT     = "22252B"
MUTED    = "606772"   # readable muted body on light
MUTED_D  = "9AA2AC"   # muted on dark
CREAM    = "F3F1EC"   # off-white text on dark

# code panel
CODEBG   = "1C2026"
C_DEF    = "E7EAEE"   # default code text
C_KEY    = "A9C3B0"   # keywords (light sage)
C_STR    = "CBB69F"   # strings (light taupe)
C_COM    = "7C8590"   # comments (slate-gray)
C_FN     = "D6C6B8"   # function names / tags
C_NUM    = "B9C7CD"

HEAD = "Arial"
BODY = "Arial"
MONO = "Courier New"

def rgb(h): return RGBColor.from_string(h)


class Deck:
    def __init__(self, w=13.333, h=7.5):
        self.prs = Presentation()
        self.prs.slide_width = Inches(w)
        self.prs.slide_height = Inches(h)
        self.W, self.H = w, h
        self.blank = self.prs.slide_layouts[6]

    def slide(self, bg=BG, textured=False, motif=True):
        s = self.prs.slides.add_slide(self.blank)
        dark = (bg == INK)
        if textured:
            import os
            p = "img/texture_dark.png"
            if os.path.exists(p):
                pic = s.shapes.add_picture(p, 0, 0, self.prs.slide_width, self.prs.slide_height)
                pic.shadow.inherit = False
            else:
                r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, self.prs.slide_width, self.prs.slide_height)
                r.fill.solid(); r.fill.fore_color.rgb = rgb(bg); r.line.fill.background(); r.shadow.inherit = False
        else:
            r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, self.prs.slide_width, self.prs.slide_height)
            r.fill.solid(); r.fill.fore_color.rgb = rgb(bg)
            r.line.fill.background(); r.shadow.inherit = False
        if motif:
            _motif(s, self.W, self.H, dark or textured)
        return s

    def save(self, path):
        self.prs.save(path)


# ---- shadow ----------------------------------------------------------------
def shadow(shape, blur=0.10, dist=0.055, direction=5400000, color=INK, alpha=20):
    """Soft outer shadow. alpha in percent of opacity (20 = subtle)."""
    spPr = shape._element.spPr
    # remove existing effectLst
    for e in spPr.findall(qn('a:effectLst')):
        spPr.remove(e)
    xml = (
        '<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f'<a:outerShdw blurRad="{int(blur*EMU_IN)}" dist="{int(dist*EMU_IN)}" '
        f'dir="{direction}" rotWithShape="0">'
        f'<a:srgbClr val="{color}"><a:alpha val="{int(alpha*1000)}"/></a:srgbClr>'
        '</a:outerShdw></a:effectLst>'
    )
    spPr.append(parse_xml(xml))


# ---- primitives ------------------------------------------------------------
def rect(slide, x, y, w, h, fill=None, line=None, lw=1.0, radius=None, shadow_=False,
         shcolor=INK, shalpha=18, dash=None):
    shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shp_type, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if radius is not None:
        try: s.adjustments[0] = radius
        except Exception: pass
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = rgb(line); s.line.width = Pt(lw)
        if dash:
            ln = s.line._get_or_add_ln()
            d = parse_xml(f'<a:prstDash xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" val="{dash}"/>')
            ln.append(d)
    if shadow_:
        shadow(s, color=shcolor, alpha=shalpha)
    return s


def _set_runs(p, runs):
    for spec in runs:
        r = p.add_run()
        r.text = spec[0]
        f = r.font
        f.size = Pt(spec[1]); f.color.rgb = rgb(spec[2])
        f.name = spec[4] if len(spec) > 4 and spec[4] else BODY
        f.bold = spec[3] if len(spec) > 3 else False
        if len(spec) > 5 and spec[5]:
            f.italic = True
        # letter spacing
        if len(spec) > 6 and spec[6]:
            rPr = r._r.get_or_add_rPr(); rPr.set('spc', str(int(spec[6])))


def text(slide, x, y, w, h, s, size=16, color=TEXT, bold=False, font=BODY,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0, italic=False,
         space_after=0.0, wrap=True, letter=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = s.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        if space_after: p.space_after = Pt(space_after)
        run = [(ln, size, color, bold, font, italic)]
        if letter: run = [(ln, size, color, bold, font, italic, letter)]
        _set_runs(p, run)
    return tb


def rich(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         spacing=1.0, wrap=True):
    """Single paragraph, multiple styled runs. runs=[(text,size,color,bold,font[,italic[,spc]]),...]"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align; p.line_spacing = spacing
    _set_runs(p, runs)
    return tb


def para(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, wrap=True):
    """Multiple paragraphs, each a list of runs. paras=[{'runs':[...],'align':,'sp':,'sa':}]"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, pd in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = pd.get('align', PP_ALIGN.LEFT)
        p.line_spacing = pd.get('sp', 1.0)
        if pd.get('sa'): p.space_after = Pt(pd['sa'])
        if pd.get('sb'): p.space_before = Pt(pd['sb'])
        _set_runs(p, pd['runs'])
    return tb


# ---- image card ------------------------------------------------------------
_ASPECT = {}
def aspect(path):
    if path not in _ASPECT:
        im = Image.open(path); _ASPECT[path] = im.width / im.height
    return _ASPECT[path]


def pic_card(slide, path, x, y, w=None, h=None, mat=True, radius=0.045,
             shadow_=True, shcolor=INK, shalpha=22, matpad=0.0):
    ar = aspect(path)
    if w and not h: h = w / ar
    if h and not w: w = h * ar
    if mat:
        m = rect(slide, x - matpad, y - matpad, w + 2*matpad, h + 2*matpad,
                 fill=SURFACE, radius=radius, shadow_=shadow_, shcolor=shcolor, shalpha=shalpha)
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    pic.shadow.inherit = False
    return pic, (x, y, w, h)


def picw(slide, path, x, y, w):
    """plain picture by width, returns (x,y,w,h)."""
    h = w / aspect(path)
    p = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    p.shadow.inherit = False
    return p, (x, y, w, h)


# ---- chips / pills ---------------------------------------------------------
def chip(slide, x, y, s, fill=SAGE_SOFT, tc=SAGE_DK, size=11.5, bold=True,
         padx=0.14, h=0.34, font=BODY, line=None):
    w = 0.115 * size/11.5 * len(s) * 0.62 + 2*padx
    r = rect(slide, x, y, w, h, fill=fill, radius=0.5, line=line, lw=1.0)
    tf = r.text_frame; tf.word_wrap = False
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    _set_runs(p, [(s, size, tc, bold, font)])
    return x + w


def arrow_down(slide, cx, y, h=0.4, color=TAUPE, w=0.26):
    a = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(cx - w/2), Inches(y), Inches(w), Inches(h))
    a.shadow.inherit = False
    a.fill.solid(); a.fill.fore_color.rgb = rgb(color); a.line.fill.background()
    return a


def arrow_right(slide, x, cy, w=0.5, color=TAUPE, h=0.26):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(cy - h/2), Inches(w), Inches(h))
    a.shadow.inherit = False
    a.fill.solid(); a.fill.fore_color.rgb = rgb(color); a.line.fill.background()
    return a


def chevron_right(slide, x, cy, size=0.22, color=TAUPE):
    a = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(cy - size/2), Inches(size*1.1), Inches(size))
    a.shadow.inherit = False
    a.fill.solid(); a.fill.fore_color.rgb = rgb(color); a.line.fill.background()
    return a


def hline(slide, x, y, w, color=MIST, weight=1.2):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x+w), Inches(y))
    ln.line.color.rgb = rgb(color); ln.line.width = Pt(weight)
    ln.shadow.inherit = False
    return ln


def vline(slide, x, y, h, color=MIST, weight=1.2):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x), Inches(y+h))
    ln.line.color.rgb = rgb(color); ln.line.width = Pt(weight)
    ln.shadow.inherit = False
    return ln


# ---- code panel ------------------------------------------------------------
def code_panel(slide, x, y, w, lines, size=12.5, title=None, lh=1.28, pad=0.26,
               header=True):
    """lines: list of paragraphs, each = list of (text,color) runs.
       Returns bottom y."""
    n = len(lines)
    body_h = n * (size/72.0) * lh * 1.2   # 1.2 = font line-height so text never clips
    head_h = 0.42 if header else 0.0
    h = head_h + body_h + 2*pad*0.72
    panel = rect(slide, x, y, w, h, fill=CODEBG, radius=0.035, shadow_=True, shcolor=INK, shalpha=30)
    if header:
        # traffic dots
        for i, c in enumerate(["E06C63", "E0B24B", "8FB98A"]):
            d = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x+pad+i*0.18), Inches(y+0.17), Inches(0.1), Inches(0.1))
            d.shadow.inherit = False; d.fill.solid(); d.fill.fore_color.rgb = rgb(c); d.line.fill.background()
        if title:
            text(slide, x+pad+0.62, y+0.10, w-2*pad-0.62, 0.28, title, size=10.5,
                 color="8A929C", font=MONO, anchor=MSO_ANCHOR.MIDDLE)
    tb = slide.shapes.add_textbox(Inches(x+pad), Inches(y+head_h+pad*0.55), Inches(w-2*pad), Inches(body_h+0.1))
    tf = tb.text_frame; tf.word_wrap = False
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, runs in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = lh
        for tval, col in runs:
            r = p.add_run(); r.text = tval
            r.font.size = Pt(size); r.font.name = MONO; r.font.color.rgb = rgb(col)
    return y + h


# ---- decorative dot motif (soft, Synora-palette) ---------------------------
# (fx, fy, radius_in, palette_key)
_DOTS = [
    (0.945, 0.05, 0.30, "accent"), (0.895, 0.145, 0.15, "soft"),
    (0.975, 0.175, 0.09, "soft2"), (0.845, 0.055, 0.07, "accent2"),
    (0.045, 0.90, 0.26, "soft"),   (0.115, 0.80, 0.12, "accent"),
    (0.02, 0.75, 0.08, "soft2"),   (0.55, 0.025, 0.06, "soft2"),
    (0.965, 0.74, 0.14, "soft"),   (0.30, 0.965, 0.09, "accent2"),
]
_PAL_LIGHT = {"soft": "E3E7EA", "soft2": "EBE1D6", "accent": "D3DDD4", "accent2": "DCC9BB"}
_PAL_DARK  = {"soft": "2E333B", "soft2": "3A342F", "accent": "3D4D44", "accent2": "4C4038"}
def _motif(slide, W, H, dark):
    pal = _PAL_DARK if dark else _PAL_LIGHT
    for fx, fy, r, key in _DOTS:
        cx, cy = fx*W, fy*H
        o = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx-r), Inches(cy-r), Inches(2*r), Inches(2*r))
        o.shadow.inherit = False
        o.fill.solid(); o.fill.fore_color.rgb = rgb(pal[key]); o.line.fill.background()


# ---- screenshot filename tag (EDUCA-style) ---------------------------------
def shot_tag(slide, x, y, name, color=INK, tc="FFFFFF", size=10):
    w = 0.072*len(name) + 0.34
    r = rect(slide, x, y, w, 0.32, fill=color, radius=0.5, shadow_=True, shalpha=24)
    tf = r.text_frame; tf.word_wrap = False
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    _set_runs(p, [(name, size, tc, True, MONO)])
    return w


# ---- monospace file/DOM tree (EDUCA-style, aligned description column) -----
def render_tree(slide, x, y, root, nodes, size=11, lh=1.36, name_gap=3,
                folder_color=None, file_color=None, desc_color=None,
                conn_color=None, root_color=None, root_desc="", w=None):
    """nodes: list of (name, desc, is_folder, children). children same shape.
       Renders one monospace block: connectors + name padded to a fixed
       column, then the description — so descriptions line up like EDUCA."""
    folder_color = folder_color or SAGE_DK
    file_color   = file_color or TEXT
    desc_color   = desc_color or SLATE
    conn_color   = conn_color or "AEB6BC"
    root_color   = root_color or TEXT
    lines = []  # (prefix, name, desc, is_folder)
    def walk(ns, prefix):
        for i, n in enumerate(ns):
            name = n[0]; desc = n[1] if len(n) > 1 else ""
            isf = n[2] if len(n) > 2 else False
            kids = n[3] if len(n) > 3 else None
            last = i == len(ns)-1
            lines.append((prefix + ("└── " if last else "├── "), name, desc, isf))
            if kids:
                walk(kids, prefix + ("    " if last else "│   "))
    walk(nodes, "")
    wchars = max([len(p)+len(nm) for p, nm, _, _ in lines] + [len(root)]) + name_gap
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w or 8.0), Inches(0.4))
    tf = tb.text_frame; tf.word_wrap = False
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    # root line
    p = tf.paragraphs[0]; p.line_spacing = lh
    rruns = [(root, size, root_color, True, MONO)]
    if root_desc:
        rruns.append((" "*(wchars-len(root)) + root_desc, size, desc_color, False, MONO))
    _set_runs(p, rruns)
    for prefix, name, desc, isf in lines:
        p = tf.add_paragraph(); p.line_spacing = lh
        runs = [(prefix, size, conn_color, False, MONO),
                (name, size, folder_color if isf else file_color, isf, MONO)]
        if desc:
            runs.append((" "*(wchars-len(prefix)-len(name)) + desc, size, desc_color, False, MONO))
        _set_runs(p, runs)
    return tb


# ---- section eyebrow (uppercase kicker + a small accent dot) ---------------
def eyebrow(slide, x, y, s, color=TAUPE, size=12):
    dd = 0.135
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y+0.035), Inches(dd), Inches(dd))
    c.shadow.inherit = False
    c.fill.solid(); c.fill.fore_color.rgb = rgb(color); c.line.fill.background()
    rich(slide, x+0.27, y, 8, 0.3, [(s.upper(), size, color, True, BODY, False, 180)])


# ---- icons (../assets/icons/<name>-<color>.png, 256px square) ---------------
ICONS = "../assets/icons/"
def icon(slide, name, x, y, size, color="ink"):
    p = slide.shapes.add_picture(ICONS + name + "-" + color + ".png",
                                 Inches(x), Inches(y), Inches(size), Inches(size))
    p.shadow.inherit = False
    return p

def icon_chip(slide, name, x, y, d=0.64, fill=SAGE, icol="white", ratio=0.52,
              shadow_=False, line=None):
    """Icon in a filled circle; (x,y) = top-left of the circle box."""
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.shadow.inherit = False
    c.fill.solid(); c.fill.fore_color.rgb = rgb(fill)
    if line is None:
        c.line.fill.background()
    else:
        c.line.color.rgb = rgb(line); c.line.width = Pt(1.2)
    if shadow_:
        shadow(c, color=INK, alpha=16)
    isz = d * ratio
    icon(slide, name, x + (d-isz)/2, y + (d-isz)/2, isz, icol)
    return c


def _tri_right(slide, cx, cy, size, color):
    a = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(cx-size/2), Inches(cy-size/2), Inches(size), Inches(size))
    a.rotation = 90; a.shadow.inherit = False
    a.fill.solid(); a.fill.fore_color.rgb = rgb(color); a.line.fill.background()
    return a


def flow_line(slide, steps, x, y, w, color=SAGE, node_d=0.74, label_size=12.5,
              sub_size=10.5, dark=False, rail_col=None, alt=None):
    """Metro-style horizontal flow: a rail, big icon nodes, arrowheads, labels below.
       steps = [(icon, title)] or [(icon, title, sub)]. `alt` = second accent to alternate."""
    n = len(steps); cy = y + node_d/2
    fx = x + node_d/2; lx = x + w - node_d/2
    span = lx - fx
    spacing = span/(n-1) if n > 1 else 0
    rail_col = rail_col or (SAGE_SOFT if not dark else INK_SOFT)
    rect(slide, fx-0.05, cy-0.05, span+0.1, 0.1, fill=rail_col, radius=0.5)
    pos = [fx + i*spacing for i in range(n)]
    for i in range(n-1):
        _tri_right(slide, (pos[i]+pos[i+1])/2, cy, 0.22, color if not alt else (color if i % 2 == 0 else alt))
    lab_col = TEXT if not dark else CREAM
    sub_col = MUTED if not dark else MUTED_D
    lab_w = min(2.3, spacing-0.12) if n > 1 else 2.3
    for i, step in enumerate(steps):
        cx = pos[i]; c = color if (alt is None or i % 2 == 0) else alt
        icon_chip(slide, step[0], cx-node_d/2, cy-node_d/2, d=node_d, fill=c, icol="white", shadow_=True)
        text(slide, cx-lab_w/2, cy+node_d/2+0.13, lab_w, 0.35, step[1], size=label_size,
             color=lab_col, bold=True, align=PP_ALIGN.CENTER, font=HEAD)
        if len(step) > 2:
            text(slide, cx-lab_w/2, cy+node_d/2+0.46, lab_w, 0.5, step[2], size=sub_size,
                 color=sub_col, align=PP_ALIGN.CENTER, spacing=1.05)
    return cy
