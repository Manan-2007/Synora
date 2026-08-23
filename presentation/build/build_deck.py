#!/usr/bin/env python3
# Synora — Web Development Project Presentation
import os
from deckkit import *
from pptx.enum.text import PP_ALIGN as AL, MSO_ANCHOR as AN

IMG = "img/"
LOGO = "../assets/"
OUT = "../Synora_Web_Development_Project_Presentation.pptx"

d = Deck()
ML, MR = 0.7, 0.7
CW = d.W - ML - MR

def footer(s, n, dark=False):
    c = MUTED_D if dark else SLATE
    text(s, ML, 7.06, 5, 0.3, "Synora  ·  Web Development Project", size=9.5, color=c)
    text(s, d.W-MR-2.0, 7.06, 2.0, 0.3, f"{n:02d} / 20", size=9.5, color=c, align=AL.RIGHT)

def head(s, kicker, title, tx=ML, tw=CW, ty=0.62):
    eyebrow(s, tx, ty, kicker)
    text(s, tx, ty+0.30, tw, 0.9, title, size=30, color=TEXT, bold=True, font=HEAD, spacing=1.0)

# =====================================================================
# 1 · TITLE
# =====================================================================
def s1():
    s = d.slide(INK)
    picw(s, LOGO+"logo-lockup-dark.png", ML, 0.62, 2.15)
    para(s, ML, 2.35, 6.4, 3.2, [
        {'runs':[("One ",44,CREAM,True,HEAD),("calm",44,SAGE_LT,True,HEAD),(" home",44,CREAM,True,HEAD)],'sp':1.02},
        {'runs':[("for everything",44,CREAM,True,HEAD)],'sp':1.02},
        {'runs':[("you're studying.",44,CREAM,True,HEAD)],'sp':1.02},
    ])
    text(s, ML, 4.86, 6.0, 0.9,
         "A student productivity & academic management workspace —\nbuilt from the HTML5, CSS3 & JavaScript syllabus.",
         size=14.5, color=MUTED_D, spacing=1.28)
    x = ML
    for lbl in ["HTML5", "CSS3", "JavaScript", "DOM", "Events"]:
        x = chip(s, x, 5.95, lbl, fill=INK_PANEL, tc=CREAM, line=SAGE, size=12, h=0.4, padx=0.17) + 0.14
    # hero
    pic_card(s, IMG+"dashboard_desktop.png", 7.35, 1.95, w=5.35, shcolor="000000", shalpha=34)
    text(s, 7.35, 5.62, 5.35, 0.3, "The Synora dashboard — the product this deck is about.",
         size=10.5, color=MUTED_D, italic=True)
    s.notes_slide.notes_text_frame.text = (
        "Open with the product, not the theory. Synora is a real, working student-productivity web app "
        "built only with HTML5, CSS3 and vanilla JavaScript — no frameworks. This deck shows how the "
        "syllabus concepts became this product, and what we built beyond them.")

# =====================================================================
# 2 · WHAT IS SYNORA
# =====================================================================
def s2():
    s = d.slide(BG)
    head(s, "The product", "What is Synora?")
    text(s, ML, 1.72, 5.3, 2.2,
         "Synora brings everything a student juggles into one calm, personal workspace — "
         "so time goes to studying, not to organising eight different apps.",
         size=15, color=MUTED, spacing=1.34)
    # feature chips grid
    feats = ["Tasks","Calendar","Quick Notes","Timetable","Attendance","CGPA","Achievements","Notifications"]
    cx, cy = ML, 3.55
    for i, f in enumerate(feats):
        col = i % 2; row = i // 2
        chip(s, ML + col*2.55, 3.55 + row*0.62, f, fill=SURFACE, tc=TEXT, line=MIST, size=12.5, h=0.46, padx=0.2)
    text(s, ML, 6.35, 5.3, 0.5, "One account.  One quiet home for the whole semester.",
         size=12.5, color=SAGE, bold=True)
    pic_card(s, IMG+"dashboard_desktop.png", 6.35, 1.5, w=6.3)
    text(s, 6.35, 5.62, 6.3, 0.3, "Dashboard — today's tasks, weekly progress and a mini-calendar at a glance.",
         size=10.5, color=SLATE, italic=True)
    footer(s, 2)
    s.notes_slide.notes_text_frame.text = (
        "Synora is a single workspace for a student's whole semester: tasks, calendar, notes, timetable, "
        "attendance, CGPA, achievements and notifications. Everything is per-user and saved in the browser. "
        "Viva: it's a single-page-per-feature app, all vanilla JS.")

# =====================================================================
# 3 · WHY WE BUILT IT
# =====================================================================
def s3():
    s = d.slide(BG)
    head(s, "The problem", "Why we built it")
    text(s, ML, 1.7, 5.4, 1.7,
         "A student's semester is scattered across a to-do app, the class WhatsApp, a notes app, "
         "a spreadsheet for attendance and another for grades. Nothing talks to anything.",
         size=14.5, color=MUTED, spacing=1.34)
    # scattered chips (left)
    scattered = ["To-do app","Sticky notes","Spreadsheet","Group chat","Calendar","Reminders"]
    bx, by = ML, 3.55
    for i, sc in enumerate(scattered):
        col = i % 2; row = i // 2
        rect(s, bx+col*2.45, by+row*0.72, 2.25, 0.56, fill=SURFACE, line=MIST, lw=1.2, radius=0.14)
        text(s, bx+col*2.45, by+row*0.72, 2.25, 0.56, sc, size=12, color=MUTED, align=AL.CENTER, anchor=AN.MIDDLE)
    text(s, ML, 5.95, 5.2, 0.4, "Scattered tools, constant context-switching", size=12, color=SLATE, italic=True)
    # arrow to synora
    chevron_right(s, 6.35, 3.9, size=0.5, color=TAUPE)
    # right: one workspace
    rect(s, 7.35, 2.05, 5.3, 3.9, fill=INK, radius=0.05, shadow_=True, shcolor=INK, shalpha=26)
    picw(s, LOGO+"logo-mark-dark.png", 7.75, 2.5, 0.72)
    text(s, 8.62, 2.52, 3.6, 0.7, "Synora", size=24, color=CREAM, bold=True, font=HEAD, anchor=AN.MIDDLE)
    text(s, 7.75, 3.5, 4.5, 0.4, "ONE CALM WORKSPACE", size=11, color=TAUPE_LT, bold=True, letter=180)
    for i, line in enumerate(["Tasks, deadlines & notes in one place",
                              "Attendance and CGPA that update live",
                              "Timetable, calendar and reminders together",
                              "Personal to each student, saved on-device"]):
        text(s, 7.9, 3.95+i*0.44, 4.6, 0.4, "—  "+line, size=12.5, color=CREAM, spacing=1.1)
    footer(s, 3)
    s.notes_slide.notes_text_frame.text = (
        "The problem is fragmentation, not a lack of tools. Students already use six apps that don't share "
        "anything. Our approach: pull the academic essentials into one quiet, unified workspace. "
        "This framing sets up why a single well-structured web app is the right shape.")

# =====================================================================
# 4 · HOW IT'S BUILT / TECH STACK + THESIS
# =====================================================================
def s4():
    s = d.slide(BG)
    head(s, "The foundation", "How Synora is built")
    text(s, ML, 1.66, CW, 0.5,
         "The syllabus building blocks are the whole foundation of the product — each layer does one job.",
         size=14.5, color=MUTED)
    rows = [
        ("HTML5", "Structure & meaning", "Semantic pages, forms, accessible controls", SAGE),
        ("CSS3", "Design & responsive layout", "Design tokens, Flexbox, Grid, media queries", TAUPE),
        ("JavaScript", "Logic & interactivity", "State, functions, arrays, objects, loops", SAGE),
        ("DOM APIs", "A dynamic interface", "Select, read, write, create & remove nodes", TAUPE),
        ("Events", "Response to the user", "addEventListener, delegation, validation", SAGE),
        ("localStorage", "Persistence", "Every user's data saved in the browser", SLATE),
    ]
    y0 = 2.4; rh = 0.66; gap = 0.12
    for i, (a, b, c, col) in enumerate(rows):
        y = y0 + i*(rh+gap)
        rect(s, ML, y, 7.7, rh, fill=SURFACE, line=MIST, lw=1.2, radius=0.09, shadow_=True, shalpha=8)
        rect(s, ML+0.14, y+0.11, 1.62, rh-0.22, fill=col, radius=0.16)
        text(s, ML+0.14, y+0.11, 1.62, rh-0.22, a, size=13, color="FFFFFF", bold=True, align=AL.CENTER, anchor=AN.MIDDLE)
        text(s, ML+2.0, y+0.09, 3.0, rh, b, size=13.5, color=TEXT, bold=True, anchor=AN.MIDDLE)
        text(s, ML+4.35, y+0.09, 3.3, rh, c, size=11.5, color=MUTED, anchor=AN.MIDDLE, spacing=1.05)
    # thesis panel
    rect(s, 8.75, 2.4, 3.9, 4.28, fill=INK, radius=0.05, shadow_=True, shcolor=INK, shalpha=26)
    text(s, 9.05, 2.72, 3.35, 0.4, "OUR THESIS", size=11, color=TAUPE_LT, bold=True, letter=200)
    text(s, 9.05, 3.2, 3.35, 3.2,
         "The syllabus gave us the web-development building blocks.\n\n"
         "Synora combines those concepts into one complete, responsive student-productivity product.",
         size=15.5, color=CREAM, spacing=1.32)
    footer(s, 4)
    s.notes_slide.notes_text_frame.text = (
        "Frame the stack as layers, each with a single responsibility. This is the deck's thesis slide: "
        "we didn't learn the syllabus separately from building Synora — the syllabus IS the foundation, and "
        "Synora is what we assembled on top. localStorage is genuinely used (16 call-sites) for persistence.")

# =====================================================================
# 5 · RELATIONSHIP DIAGRAM
# =====================================================================
def s5():
    s = d.slide(INK)
    eyebrow(s, ML, 0.62, "How the layers work together", color=TAUPE_LT)
    text(s, ML, 0.92, CW, 0.7, "Five questions, one product", size=30, color=CREAM, bold=True, font=HEAD)
    cols = [
        ("HTML5", "What exists?", "Structure & content", SAGE_LT),
        ("CSS3", "How does it look?", "Design & layout", TAUPE_LT),
        ("JavaScript", "How does it behave?", "Logic & state", SAGE_LT),
        ("DOM", "How does JS change\nthe page?", "Live updates", TAUPE_LT),
        ("Events", "How do actions\ntrigger behaviour?", "User → response", SAGE_LT),
    ]
    n = len(cols); cw = 2.16; gap = (CW - n*cw) / (n-1)
    y = 2.15; ch = 3.15
    for i, (t, q, sub, ac) in enumerate(cols):
        x = ML + i*(cw+gap)
        rect(s, x, y, cw, ch, fill=INK_PANEL, radius=0.07, line=INK_SOFT, lw=1.0, shadow_=True, shcolor="000000", shalpha=30)
        text(s, x, y+0.16, cw, 0.5, t, size=16.5, color=ac, bold=True, align=AL.CENTER, font=HEAD)
        hline(s, x+0.35, y+0.78, cw-0.7, color=INK_SOFT, weight=1.0)
        text(s, x+0.16, y+0.95, cw-0.32, 1.3, q, size=14.5, color=CREAM, bold=True, align=AL.CENTER, spacing=1.08, anchor=AN.TOP)
        text(s, x+0.16, y+ch-0.62, cw-0.32, 0.5, sub, size=11, color=MUTED_D, align=AL.CENTER, anchor=AN.MIDDLE)
        if i < n-1:
            chevron_right(s, x+cw+gap/2-0.12, y+ch/2, size=0.26, color=TAUPE)
    # result bar
    rect(s, ML, 5.85, CW, 0.86, fill=SAGE, radius=0.09, shadow_=True, shcolor="000000", shalpha=28)
    rich(s, ML, 5.85, CW, 0.86, [
        ("The five layers combine into  ", 15, "EAF0EA", False, BODY),
        ("Synora", 17, "FFFFFF", True, HEAD),
        ("  —  a student productivity workspace.", 15, "EAF0EA", False, BODY),
    ], align=AL.CENTER, anchor=AN.MIDDLE)
    footer(s, 5, dark=True)
    s.notes_slide.notes_text_frame.text = (
        "This is the key conceptual diagram. Each syllabus layer answers one question: HTML = what exists, "
        "CSS = how it looks, JS = how it behaves, DOM = how JS changes the page, Events = how user actions "
        "trigger behaviour. Together they produce the product. Likely viva question: 'how do HTML, CSS and JS "
        "relate?' — answer with this line-up.")

# =====================================================================
# 6 · HTML5 STRUCTURE  [SNIPPET 1]
# =====================================================================
def s6():
    s = d.slide(BG)
    head(s, "Lecture 1–6  ·  HTML5", "Building the structure")
    text(s, ML, 1.66, 5.55, 1.5,
         "Every Synora page is written with semantic HTML5 — elements that describe what each region "
         "means, not just how it looks. The browser, assistive tech and our own code all read that structure.",
         size=13.5, color=MUTED, spacing=1.3)
    # tag list
    tags = [("<header>",""),("<nav>","aria-label"),("<main>","×14"),("<section>","×11 / aria-labelledby"),
            ("<aside>","×12 sidebar"),("<footer>",""),("<figure>","×3")]
    for i,(t,note) in enumerate(tags):
        y = 3.35 + i*0.44
        rich(s, ML, y, 5.5, 0.4, [(t, 13, SAGE_DK, True, MONO), (("   "+note) if note else "", 11, SLATE, False, BODY)])
    code_panel(s, 6.55, 1.6, 6.12, [
        [("<header ", C_FN), ("class", C_KEY), ("=", C_DEF), ('"site-header"', C_STR), (">", C_FN)],
        [("  <nav ", C_FN), ("aria-label", C_KEY), ("=", C_DEF), ('"Primary"', C_STR), ("> … </nav>", C_FN)],
        [("</header>", C_FN)],
        [("", C_DEF)],
        [("<main ", C_FN), ("id", C_KEY), ("=", C_DEF), ('"main"', C_STR), (">", C_FN)],
        [("  <section ", C_FN), ("class", C_KEY), ("=", C_DEF), ('"hero"', C_STR), (">…</section>", C_FN)],
        [("  <section ", C_FN), ("aria-labelledby", C_KEY), ("=", C_DEF), ('"features"', C_STR), (">…", C_FN)],
        [("  </section>", C_FN)],
        [("</main>", C_FN)],
        [("", C_DEF)],
        [("<footer ", C_FN), ("class", C_KEY), ("=", C_DEF), ('"site-footer"', C_STR), (">…</footer>", C_FN)],
    ], size=12.5, title="index.html — landing structure")
    text(s, 6.55, 6.75, 6.12, 0.3, "Real markup from index.html — trimmed.", size=10, color=SLATE, italic=True)
    footer(s, 6)
    s.notes_slide.notes_text_frame.text = (
        "Semantic tags carry meaning. header/nav/main/section/aside/footer appear across every page — main "
        "14×, section 11×, aside 12×. Sections use aria-labelledby to tie a heading to its region. "
        "Viva: why semantic HTML? Accessibility, SEO, and maintainable structure our JS can target reliably.")

# =====================================================================
# 7 · SEMANTIC + ACCESSIBILITY
# =====================================================================
def s7():
    s = d.slide(BG)
    head(s, "HTML5  ·  Accessibility basics", "Meaningful, accessible controls")
    pic_card(s, IMG+"settings.png", ML, 1.6, w=6.35)
    text(s, ML, 5.72, 6.35, 0.3, "Settings — every field has a real <label> and input type.",
         size=10.5, color=SLATE, italic=True)
    # callouts right
    items = [
        ("Labels for every input", "16 <label> elements — clicking a label focuses its field."),
        ("ARIA where it helps", "91 aria-* attributes: aria-label, aria-labelledby, aria-live regions."),
        ("Right input types", "type=email, number, date, password — the keyboard and validation adapt."),
        ("Alt text on imagery", "18 alt attributes; decorative art is marked aria-hidden."),
        ("Keyboard-friendly", "A skip-link and a .visually-hidden pattern for screen-reader-only text."),
    ]
    y = 1.66
    for t, b in items:
        rect(s, 7.35, y, 0.13, 0.86, fill=SAGE, radius=0.4)
        text(s, 7.66, y, 5.0, 0.4, t, size=13.5, color=TEXT, bold=True)
        text(s, 7.66, y+0.34, 5.0, 0.6, b, size=11, color=MUTED, spacing=1.16)
        y += 1.02
    footer(s, 7)
    s.notes_slide.notes_text_frame.text = (
        "Be honest: this is accessibility BASICS, not certified WCAG compliance. What's genuinely there: "
        "labels (16), aria-* (91), correct input types, alt text (18), a skip-link and a visually-hidden "
        "utility. Viva: what makes a form accessible? Programmatic label-to-input association and correct "
        "input types.")

# =====================================================================
# 8 · CSS3 DESIGN SYSTEM
# =====================================================================
def s8():
    s = d.slide(BG)
    head(s, "Lecture 1–6  ·  CSS3", "One design system, not many pages")
    text(s, ML, 1.66, 6.0, 1.0,
         "CSS3 custom properties define the look once — ~142 design tokens for colour, type, spacing, radius "
         "and shadow. Every card, button and form reuses them, so the whole app feels like one product.",
         size=13.5, color=MUTED, spacing=1.28)
    # palette swatches
    pal = [("Ink","22252B"),("Sage","53655C"),("Slate","A3AEB1"),("Mist","E5E8EB"),("Taupe","B09F95")]
    sw = 1.5; gap = 0.18; x0 = ML
    for i,(name,hexv) in enumerate(pal):
        x = x0 + i*(sw+gap)
        rect(s, x, 3.15, sw, 1.05, fill=hexv, radius=0.1, line=MIST, lw=1.0, shadow_=True, shalpha=10)
        text(s, x, 4.28, sw, 0.3, name, size=12, color=TEXT, bold=True, align=AL.CENTER)
        text(s, x, 4.55, sw, 0.3, "#"+hexv, size=10, color=SLATE, align=AL.CENTER, font=MONO)
    # type + tokens
    rect(s, ML, 5.2, 4.4, 1.5, fill=SURFACE, line=MIST, lw=1.2, radius=0.09, shadow_=True, shalpha=8)
    text(s, ML+0.3, 5.38, 3.9, 0.4, "Typography", size=11, color=TAUPE, bold=True, letter=140)
    text(s, ML+0.3, 5.72, 3.9, 0.5, "Manrope", size=20, color=TEXT, bold=True, font=HEAD)
    text(s, ML+0.3, 6.16, 3.9, 0.4, "headings  ·  Inter for body text", size=12, color=MUTED)
    rect(s, ML+4.7, 5.2, 4.4, 1.5, fill=SURFACE, line=MIST, lw=1.2, radius=0.09, shadow_=True, shalpha=8)
    text(s, ML+5.0, 5.38, 3.9, 0.4, "Reused components", size=11, color=TAUPE, bold=True, letter=140)
    text(s, ML+5.0, 5.74, 3.9, 0.5, ".btn   .card   .field   .chip", size=15, color=SAGE_DK, bold=True, font=MONO)
    text(s, ML+5.0, 6.16, 3.9, 0.4, "defined once, used everywhere", size=12, color=MUTED)
    # right screenshot
    pic_card(s, IMG+"notes.png", 9.35, 2.0, w=3.35)
    text(s, 9.35, 4.16, 3.35, 0.3, "Quick Notes — same cards,\nbuttons and spacing tokens.",
         size=10, color=SLATE, italic=True, spacing=1.1)
    footer(s, 8)
    s.notes_slide.notes_text_frame.text = (
        "CSS3 was used to build a design SYSTEM, not to style each page by hand. ~142 custom properties in "
        "tokens.css hold the palette, type scale, spacing, radii and shadows; light/dark themes swap the same "
        "tokens. Viva: why CSS variables? Single source of truth — change a token, the whole app updates.")

# =====================================================================
# 9 · BOX MODEL  [SNIPPET 2]
# =====================================================================
def s9():
    s = d.slide(BG)
    head(s, "CSS3  ·  The Box Model", "Every component is a box")
    # box model diagram (nested) — left
    bx, by, bw, bh = ML, 1.78, 4.5, 2.95
    rect(s, bx, by, bw, bh, fill="EFE7DF", radius=0.03)  # margin
    text(s, bx+0.12, by+0.07, 2, 0.3, "margin", size=10, color=TAUPE, bold=True)
    rect(s, bx+0.5, by+0.4, bw-1.0, bh-0.8, fill="DCE3DC", radius=0.03)  # border
    text(s, bx+0.62, by+0.46, 2, 0.3, "border", size=10, color=SAGE_DK, bold=True)
    rect(s, bx+0.95, by+0.78, bw-1.9, bh-1.56, fill="E9EDF0", radius=0.03)  # padding
    text(s, bx+1.07, by+0.84, 2, 0.3, "padding", size=10, color=SLATE, bold=True)
    rect(s, bx+1.5, by+1.16, bw-3.0, bh-2.32, fill=SURFACE, radius=0.05, line=MIST, lw=1.0)
    text(s, bx+1.5, by+1.16, bw-3.0, bh-2.32, "content", size=12, color=TEXT, bold=True, align=AL.CENTER, anchor=AN.MIDDLE)
    text(s, bx, by+bh+0.12, bw, 0.3, "content → padding → border → margin", size=11.5, color=MUTED, align=AL.CENTER)
    # snippet — right
    code_panel(s, 5.9, 1.78, 6.75, [
        [(".stat-card", C_FN), (" {", C_DEF)],
        [("  padding", C_KEY), (": ", C_DEF), ("20px 24px", C_NUM), (";", C_DEF), ("      /* space inside */", C_COM)],
        [("  background", C_KEY), (": ", C_DEF), ("var(--color-surface)", C_STR), (";", C_DEF)],
        [("  border", C_KEY), (": ", C_DEF), ("1px solid ", C_NUM), ("var(--color-border)", C_STR), (";", C_DEF)],
        [("  border-radius", C_KEY), (": ", C_DEF), ("var(--radius-lg)", C_STR), (";", C_DEF)],
        [("  box-shadow", C_KEY), (": ", C_DEF), ("var(--shadow-sm)", C_STR), (";", C_DEF)],
        [("}", C_DEF)],
    ], size=12.5, title="styles/app.css — a dashboard stat card")
    # full-width strip: the real row of boxes
    text(s, ML, 5.22, CW, 0.3, "Those same tokens build every card — here, the CGPA summary row:",
         size=12, color=TEXT, bold=True)
    stw = 9.0
    pic_card(s, IMG+"cgpa_cards.png", (d.W-stw)/2, 5.58, w=stw)
    footer(s, 9)
    s.notes_slide.notes_text_frame.text = (
        "The box model is not theory here — every card IS a box. .stat-card sets padding (space inside), a "
        "1px border, rounded corners and a soft shadow, all via tokens. Viva: name the four box-model layers "
        "from inside out — content, padding, border, margin.")

# =====================================================================
# 10 · FLEXBOX + GRID
# =====================================================================
def s10():
    s = d.slide(BG)
    head(s, "CSS3  ·  Layout", "Flexbox for rows, Grid for structure")
    # left flexbox
    text(s, ML, 1.72, 5.6, 0.4, "Flexbox — one-dimensional rows", size=15, color=SAGE_DK, bold=True)
    pic_card(s, IMG+"tasks.png", ML, 2.2, w=5.6)
    rich(s, ML, 4.35, 5.6, 0.6, [("display: flex", 12.5, SAGE_DK, True, MONO),
                                  ("  aligns the top bar, toolbars and button rows.", 11.5, MUTED, False, BODY)])
    text(s, ML, 4.72, 5.6, 0.5, "Used in ~116 rules — align-items, gap, flex-wrap.", size=11, color=SLATE)
    # right grid
    gx = 6.95
    text(s, gx, 1.72, 5.6, 0.4, "Grid — two-dimensional layouts", size=15, color=TAUPE, bold=True)
    pic_card(s, IMG+"calendar.png", gx, 2.2, w=5.7)
    rich(s, gx, 4.35, 5.7, 0.6, [("display: grid", 12.5, TAUPE, True, MONO),
                                 ("  builds the calendar, dashboard & note grids.", 11.5, MUTED, False, BODY)])
    text(s, gx, 4.72, 5.7, 0.5, "Used in ~51 rules — repeat(7, minmax(0,1fr)), gap.", size=11, color=SLATE)
    # bottom code strip
    code_panel(s, ML, 5.5, CW, [
        [(".cal-grid", C_FN), (" { ", C_DEF), ("display", C_KEY), (": ", C_DEF), ("grid", C_NUM),
         ("; ", C_DEF), ("grid-template-columns", C_KEY), (": ", C_DEF), ("repeat(7, minmax(0, 1fr))", C_STR),
         ("; ", C_DEF), ("gap", C_KEY), (": ", C_DEF), ("6px", C_NUM), ("; }", C_DEF)],
    ], size=12.5, title="styles/app.css — the seven-column calendar grid", header=True)
    footer(s, 10)
    s.notes_slide.notes_text_frame.text = (
        "Flexbox and Grid do different jobs. Flexbox (116 rules) lays out one-dimensional rows — the sticky "
        "top bar, toolbars, button groups. Grid (51 rules) handles two-dimensional structure — the calendar's "
        "7 columns, the dashboard cards, the notes masonry. Viva: why Grid for the dashboard? It's a real 2-D "
        "arrangement of rows and columns.")

# =====================================================================
# 11 · RESPONSIVE (MANDATORY 3-DEVICE)
# =====================================================================
def s11():
    s = d.slide(INK)
    eyebrow(s, ML, 0.6, "CSS3  ·  Responsive design", color=TAUPE_LT)
    text(s, ML, 0.9, CW, 0.6, "One page, every screen", size=30, color=CREAM, bold=True, font=HEAD)
    # three devices, standing on a common baseline so the labels align
    BOT = 5.82
    def dev(img, x, w, lx, lw, label):
        h = w / aspect(IMG+img)
        pic_card(s, IMG+img, x, BOT-h, w=w, shcolor="000000", shalpha=40)
        text(s, lx, BOT+0.14, lw, 0.3, label, size=12, color=SAGE_LT, bold=True, align=AL.CENTER)
    dev("resp_desktop.png", ML, 6.15, ML, 6.15, "1440px  ·  Desktop")
    dev("resp_tablet.png", 7.7, 2.35, 7.7, 2.35, "768px  ·  Tablet")
    dev("resp_mobile.png", 10.35, 1.6, 9.95, 2.4, "390px  ·  Mobile")
    # strategy strip
    rect(s, ML, 6.4, CW, 0.6, fill=INK_PANEL, radius=0.09, line=INK_SOFT, lw=1.0)
    rich(s, ML+0.3, 6.4, CW-0.6, 0.6, [
        ("Desktop-first  ", 13, TAUPE_LT, True, BODY),
        ("— Flexbox + Grid reflow, ", 12.5, CREAM, False, BODY),
        ("max-width", 12.5, SAGE_LT, True, MONO),
        (" media queries adapt down, and the nav collapses to a CSS-only menu.", 12.5, CREAM, False, BODY),
    ], anchor=AN.MIDDLE)
    footer(s, 11, dark=True)
    s.notes_slide.notes_text_frame.text = (
        "This proves responsiveness rather than claiming it — the SAME landing page at 1440/768/390. Note the "
        "grid hero collapses to one column, buttons go full-width, and the navbar becomes a hamburger. Be "
        "accurate: the strategy is DESKTOP-FIRST (max-width breakpoints adapting down), not mobile-first. The "
        "menu toggle is pure CSS — a hidden checkbox, no JavaScript.")

# =====================================================================
# 12 · JAVASCRIPT FUNDAMENTALS  [SNIPPET 3]
# =====================================================================
def s12():
    s = d.slide(BG)
    head(s, "Lecture 7–12  ·  JavaScript", "From a static page to a product")
    text(s, ML, 1.66, 5.7, 1.0,
         "HTML and CSS give structure and style; JavaScript adds the behaviour. The five fundamentals map "
         "directly onto real parts of Synora:",
         size=13.5, color=MUTED, spacing=1.28)
    rows = [
        ("Variables", "hold application state", "state = { view, sort, query }"),
        ("Functions", "reusable operations", "render()  ·  toggleComplete()"),
        ("Arrays", "lists of things", "tasks  ·  subjects  ·  timetable"),
        ("Objects", "structured records", "{ id, title, priority, deadline }"),
        ("Loops", "render each item", ".map()  ·  .filter()  ·  .forEach()"),
    ]
    y = 2.9
    for a,b,c in rows:
        rect(s, ML, y, 5.85, 0.62, fill=SURFACE, line=MIST, lw=1.2, radius=0.09, shadow_=True, shalpha=8)
        text(s, ML+0.22, y, 1.55, 0.62, a, size=13.5, color=SAGE_DK, bold=True, anchor=AN.MIDDLE)
        text(s, ML+1.8, y, 2.0, 0.62, b, size=11.5, color=MUTED, anchor=AN.MIDDLE)
        text(s, ML+3.75, y, 2.0, 0.62, c, size=10.5, color=TEXT, font=MONO, anchor=AN.MIDDLE)
        y += 0.72
    code_panel(s, 6.85, 1.62, 5.82, [
        [("function", C_KEY), (" filtered() {", C_DEF)],
        [("  var", C_KEY), (" items = all().", C_DEF), ("slice", C_FN), ("();", C_DEF)],
        [("  items = items.", C_DEF), ("filter", C_FN), ("(", C_DEF), ("function", C_KEY), ("(t) {", C_DEF)],
        [("    return", C_KEY), (" t.title.", C_DEF), ("indexOf", C_FN), ("(q) >= ", C_DEF), ("0", C_NUM), (";", C_DEF)],
        [("  });", C_DEF)],
        [("  items.", C_DEF), ("sort", C_FN), ("(byDeadline);", C_DEF)],
        [("  return", C_KEY), (" items;", C_DEF)],
        [("}", C_DEF)],
    ], size=12.5, title="scripts/tasks.js — arrays, objects & loops")
    text(s, 6.85, 5.35, 5.82, 1.4,
         "One function turns an array of task objects into the filtered, sorted list you see — "
         "variables, functions, arrays, objects and loops all in a few lines.",
         size=12, color=MUTED, spacing=1.28)
    footer(s, 12)
    s.notes_slide.notes_text_frame.text = (
        "Connect each fundamental to something concrete. State lives in variables; render() and "
        "toggleComplete() are reusable functions; tasks/subjects are arrays; each task is an object; .map / "
        ".filter / .forEach are the loops. The snippet shows all five working together. Viva: difference "
        "between an array and an object? Ordered list vs keyed record.")

# =====================================================================
# 13 · DOM MANIPULATION  [SNIPPET 4]
# =====================================================================
def s13():
    s = d.slide(BG)
    head(s, "Lecture 23–26  ·  DOM", "Making the interface dynamic")
    # DOM tree (left)
    text(s, ML, 1.66, 5.6, 0.4, "The page is a tree of nodes", size=14.5, color=TEXT, bold=True)
    tree = [
        (0,"document"),(1,"main.page"),(2,"section.toolbar"),(3,"input[data-search]"),
        (2,"div[data-list]"),(3,"div.tlist"),(4,"div.trow  ← one per task"),
    ]
    for i,(depth,label) in enumerate(tree):
        y = 2.16 + i*0.42
        col = SAGE_DK if depth<=1 else (TAUPE if "trow" in label else MUTED)
        bold = depth<=1 or "trow" in label
        text(s, ML+depth*0.42, y, 5.4, 0.4, ("└─ " if depth>0 else "")+label, size=12.5,
             color=col, font=MONO, bold=bold)
    # selectors used
    rect(s, ML, 5.35, 5.7, 1.5, fill=SURFACE, line=MIST, lw=1.2, radius=0.09, shadow_=True, shalpha=8)
    text(s, ML+0.25, 5.5, 5.3, 0.3, "SELECTORS ACTUALLY USED", size=10.5, color=TAUPE, bold=True, letter=140)
    for i,(fn,ct) in enumerate([("querySelector()","×161"),("querySelectorAll()","×23"),("getElementById()","×9")]):
        yy = 5.85 + i*0.32
        rich(s, ML+0.25, yy, 5.3, 0.3, [(fn, 12.5, SAGE_DK, True, MONO), ("   "+ct, 11.5, SLATE, False, BODY)])
    # node lifecycle snippet (right)
    text(s, 6.85, 1.66, 5.8, 0.4, "Create, configure, insert — then remove", size=14.5, color=TEXT, bold=True)
    code_panel(s, 6.85, 2.12, 5.82, [
        [("var", C_KEY), (" el = document.", C_DEF), ("createElement", C_FN), ("(", C_DEF), ('"div"', C_STR), (");", C_DEF), ("  // 1 create", C_COM)],
        [("el.className = ", C_DEF), ('"toast"', C_STR), (";", C_DEF)],
        [("el.", C_DEF), ("setAttribute", C_FN), ("(", C_DEF), ('"role"', C_STR), (", ", C_DEF), ('"status"', C_STR), (");", C_DEF), (" // 2 config", C_COM)],
        [("el.textContent = message;", C_DEF)],
        [("region.", C_DEF), ("appendChild", C_FN), ("(el);", C_DEF), ("        // 3 insert", C_COM)],
        [("…", C_DEF)],
        [("el.", C_DEF), ("remove", C_FN), ("();", C_DEF), ("                    // 4 delete", C_COM)],
    ], size=12, title="scripts/app.js — the toast lifecycle")
    text(s, 6.85, 5.35, 5.82, 1.4,
         "JS reads and writes the page through the DOM: select an element, change its text, classes or "
         "attributes, build new nodes, and remove them when done. classList is toggled in ~35 places.",
         size=12, color=MUTED, spacing=1.28)
    footer(s, 13)
    s.notes_slide.notes_text_frame.text = (
        "DOM = the browser's live object model of the HTML. We select nodes (querySelector 161×, "
        "querySelectorAll 23×, getElementById 9×), read/write text and classes, and create/remove nodes. The "
        "toast shows the full node lifecycle: createElement → configure → appendChild → remove. Viva: what is "
        "the DOM? A tree of objects representing the document that JS can change on the fly.")

# =====================================================================
# 14 · EVENTS — CLICK TO ACTION  [SNIPPET 5]
# =====================================================================
def s14():
    s = d.slide(BG)
    head(s, "Lecture 27–30  ·  Events", "From a click to an updated screen")
    # flow
    steps = ["User acts","Event fires","addEventListener","Handler runs","DOM / data update","Screen re-renders"]
    x = ML; y = 1.72
    bw = 1.83; gap=0.15
    for i, st in enumerate(steps):
        bx = ML + i*(bw+gap)
        fill = SAGE if i in (2,) else SURFACE
        tc = "FFFFFF" if i==2 else TEXT
        rect(s, bx, y, bw, 0.72, fill=fill, line=None if i==2 else MIST, lw=1.2, radius=0.1, shadow_=True, shalpha=10)
        text(s, bx+0.08, y, bw-0.16, 0.72, st, size=11.5, color=tc, bold=True, align=AL.CENTER, anchor=AN.MIDDLE, spacing=1.0)
        if i < len(steps)-1:
            chevron_right(s, bx+bw+gap/2-0.09, y+0.36, size=0.18, color=TAUPE)
    # snippet
    code_panel(s, ML, 2.95, 6.5, [
        [("var", C_KEY), (" addBtn = document.", C_DEF), ("querySelector", C_FN), ("(", C_DEF), ('"[data-add-task]"', C_STR), (");", C_DEF)],
        [("addBtn.", C_DEF), ("addEventListener", C_FN), ("(", C_DEF), ('"click"', C_STR), (", ", C_DEF), ("function", C_KEY), ("() {", C_DEF)],
        [("  openForm(", C_DEF), ("null", C_KEY), (", render);", C_DEF), ("   // click → open the task form", C_COM)],
        [("});", C_DEF)],
        [("", C_DEF)],
        [("search.", C_DEF), ("addEventListener", C_FN), ("(", C_DEF), ('"input"', C_STR), (", ", C_DEF), ("function", C_KEY), ("(e) {", C_DEF)],
        [("  state.query = e.target.value; render();", C_DEF), (" // each keystroke", C_COM)],
        [("});", C_DEF)],
    ], size=12, title="scripts/tasks.js — click & input listeners")
    # before / after
    pic_card(s, IMG+"attendance.png", 7.4, 2.95, w=5.2)
    text(s, 7.4, 6.35, 5.2, 0.6,
         "Attendance +/– buttons fire click events that recompute every percentage instantly. "
         "addEventListener is used ~93×.",
         size=11.5, color=MUTED, spacing=1.22)
    footer(s, 14)
    s.notes_slide.notes_text_frame.text = (
        "Events are the whole interaction model. The flow: user acts → event fires → addEventListener's "
        "handler runs → it changes data/DOM → the screen re-renders. Real examples: the add-task click, the "
        "search input firing on every keystroke, attendance +/- buttons. addEventListener appears ~93 times; "
        "preventDefault ~11.")

# =====================================================================
# 15 · PROPAGATION & DELEGATION  [SNIPPET 6]
# =====================================================================
def s15():
    s = d.slide(BG)
    head(s, "Events  ·  Propagation & delegation", "One listener, many buttons")
    # bubbling diagram
    text(s, ML, 1.66, 5.6, 0.4, "Events bubble up the tree", size=14.5, color=TEXT, bold=True)
    levels=[("container  [data-list]",SAGE_DK,0),("div.trow",MUTED,0.5),("button  [data-toggle]",TAUPE,1.0)]
    for i,(lab,col,ind) in enumerate(levels):
        y=2.15+i*0.62
        rect(s, ML+ind*0.5, y, 4.6-ind, 0.5, fill=SURFACE, line=MIST, lw=1.2, radius=0.1)
        text(s, ML+ind*0.5+0.2, y, 4.2-ind, 0.5, lab, size=12, color=col, font=MONO, bold=True, anchor=AN.MIDDLE)
    # up arrows
    for i in range(2):
        a = s.shapes.add_shape(MSO_SHAPE.UP_ARROW, Inches(ML+4.9), Inches(2.35+i*0.62), Inches(0.26), Inches(0.42))
        a.shadow.inherit=False; a.fill.solid(); a.fill.fore_color.rgb=rgb(TAUPE); a.line.fill.background()
    text(s, ML+5.3, 2.4, 1.6, 0.9, "click on a\nbutton bubbles\nup to the list",
         size=10.5, color=MUTED, spacing=1.12)
    # honesty box
    rect(s, ML, 4.35, 5.7, 2.35, fill=SURFACE, line=MIST, lw=1.2, radius=0.08, shadow_=True, shalpha=8)
    text(s, ML+0.28, 4.55, 5.2, 0.3, "IN SYNORA — HONESTLY", size=10.5, color=TAUPE, bold=True, letter=140)
    used = [("Event bubbling", "used — clicks bubble to the parent"),
            ("Event delegation", "used — .closest() in ~28 handlers"),
            ("preventDefault", "used — forms & links (~11)")]
    for i,(a,b) in enumerate(used):
        yy=4.9+i*0.42
        rect(s, ML+0.28, yy+0.03, 0.16,0.16, fill=SAGE, radius=0.5)
        rich(s, ML+0.56, yy, 5.0, 0.4, [(a+"  ",12.5,TEXT,True,BODY),(b,11,MUTED,False,BODY)])
    yy=4.9+3*0.42+0.05
    rect(s, ML+0.28, yy+0.03, 0.16,0.16, fill=SLATE, radius=0.5)
    rich(s, ML+0.56, yy, 5.0, 0.4, [("Explicit capturing  ",12.5,TEXT,True,BODY),
                                    ("concept only — not a pattern we use",11,SLATE,False,BODY,True)])
    # delegation snippet
    code_panel(s, 6.95, 1.62, 5.72, [
        [("// ONE listener handles every task row", C_COM)],
        [("container.", C_DEF), ("addEventListener", C_FN), ("(", C_DEF), ('"click"', C_STR), (", ", C_DEF), ("function", C_KEY), ("(e){", C_DEF)],
        [("  var", C_KEY), (" t = e.target.", C_DEF), ("closest", C_FN), ("(", C_DEF), ('"[data-toggle]"', C_STR), (");", C_DEF)],
        [("  var", C_KEY), (" d = e.target.", C_DEF), ("closest", C_FN), ("(", C_DEF), ('"[data-del]"', C_STR), (");", C_DEF)],
        [("  if", C_KEY), (" (t) { toggleComplete(t...); rerender(); }", C_DEF)],
        [("  else if", C_KEY), (" (d) { remove(d...); rerender(); }", C_DEF)],
        [("});", C_DEF)],
    ], size=12, title="scripts/tasks.js — event delegation")
    text(s, 6.95, 4.85, 5.72, 1.8,
         "Task rows are re-rendered constantly, so binding a listener to each button would be fragile. "
         "Instead one listener sits on the container and uses e.target.closest() to find which button was "
         "clicked — that is event delegation, powered by bubbling.",
         size=12, color=MUTED, spacing=1.3)
    footer(s, 15)
    s.notes_slide.notes_text_frame.text = (
        "This is the honesty slide the brief demands. USED: bubbling, delegation via closest() (28 handlers), "
        "preventDefault. NOT a real pattern: explicit event capturing — say so plainly. Explain why "
        "delegation matters here: rows are re-rendered, so one durable listener on the parent is more robust "
        "than per-button listeners. Viva: bubbling vs capturing — bubbling goes child→parent, capturing "
        "parent→child.")

# =====================================================================
# 16 · FORMS + VALIDATION
# =====================================================================
def s16():
    s = d.slide(BG)
    head(s, "Events  ·  Forms & validation", "Catching mistakes before they save")
    pic_card(s, IMG+"login_desktop.png", ML, 1.62, w=6.05)
    text(s, ML, 5.55, 6.05, 0.3, "Sign-in and sign-up validate on submit and clear errors as you type.",
         size=10.5, color=SLATE, italic=True)
    # validation flow chips + honesty
    text(s, 7.05, 1.66, 5.6, 0.4, "Two layers of validation", size=14.5, color=TEXT, bold=True)
    layers=[("HTML attributes","required, type=email, minlength — the browser's first check"),
            ("JavaScript on submit","e.preventDefault(), trim, regex, then friendly inline errors"),
            ("Live error clearing","the error disappears the moment the field is corrected")]
    y=2.15
    for a,b in layers:
        rect(s, 7.05, y, 0.13, 0.7, fill=TAUPE, radius=0.4)
        text(s, 7.35, y, 5.3, 0.4, a, size=13, color=TEXT, bold=True)
        text(s, 7.35, y+0.32, 5.3, 0.5, b, size=11, color=MUTED, spacing=1.15)
        y+=0.86
    code_panel(s, 7.05, 4.85, 5.62, [
        [("form.", C_DEF), ("addEventListener", C_FN), ("(", C_DEF), ('"submit"', C_STR), (", ", C_DEF), ("function", C_KEY), ("(e){", C_DEF)],
        [("  e.", C_DEF), ("preventDefault", C_FN), ("();", C_DEF), ("            // stop native submit", C_COM)],
        [("  if", C_KEY), (" (!EMAIL_RE.", C_DEF), ("test", C_FN), ("(email))", C_DEF)],
        [("    setError(", C_DEF), ('"email"', C_STR), (", ", C_DEF), ('"Check your email."', C_STR), (");", C_DEF)],
        [("});", C_DEF)],
    ], size=11.5, title="scripts/auth.js — sign-up validation")
    footer(s, 16)
    s.notes_slide.notes_text_frame.text = (
        "Validation is layered. HTML attributes (required, type=email, minlength) are the first gate. On "
        "submit, JS calls preventDefault(), trims input, tests it against regexes, and shows friendly inline "
        "errors instead of the browser default. Errors clear live as the user fixes the field. Viva: why "
        "validate in JS if HTML already does? Custom messages, cross-field checks like password confirmation, "
        "and control over UX.")

# =====================================================================
# 17 · LOADING EXPERIENCE
# =====================================================================
def s17():
    s = d.slide(INK)
    eyebrow(s, ML, 0.6, "User experience  ·  the loading screen", color=TAUPE_LT)
    text(s, ML, 0.9, CW, 0.6, "Before you see Synora", size=30, color=CREAM, bold=True, font=HEAD)
    # flow of three states
    stages=[("Splash","logo, ring & wordmark\nanimate in"),
            ("Handover","splash fades & lifts as the\npage arrives underneath"),
            ("Landing","the hero's staggered\nentrance begins")]
    bw=3.55; gap=0.55; y=1.85
    for i,(t,b) in enumerate(stages):
        x=ML+i*(bw+gap)
        rect(s, x, y, bw, 1.5, fill=INK_PANEL, radius=0.08, line=INK_SOFT, lw=1.0, shadow_=True, shcolor="000000", shalpha=28)
        text(s, x+0.3, y+0.22, bw-0.6, 0.4, t, size=16, color=SAGE_LT, bold=True, font=HEAD)
        text(s, x+0.3, y+0.66, bw-0.6, 0.7, b, size=11.5, color=CREAM, spacing=1.2)
        if i<2:
            chevron_right(s, x+bw+gap/2-0.11, y+0.75, size=0.28, color=TAUPE)
    # mechanism (honest)
    rect(s, ML, 3.75, 6.1, 2.95, fill=INK_PANEL, radius=0.07, line=INK_SOFT, lw=1.0)
    text(s, ML+0.32, 3.98, 5.5, 0.3, "HOW IT ACTUALLY WORKS", size=10.5, color=TAUPE_LT, bold=True, letter=160)
    mech=[("CSS @keyframes","spin, breathe, pop & sheen — pure CSS animation"),
          ("A CSS transition","the splash fades + lifts on the .is-hidden class"),
          ("JavaScript timing","shows for ≥2000ms, waits for DOMContentLoaded"),
          ("DOM + state","toggles classes, then loader.remove() clears the node"),
          ("Reduced motion","honoured — animation off, wait cut to 600ms")]
    for i,(a,b) in enumerate(mech):
        yy=4.35+i*0.46
        rich(s, ML+0.32, yy, 5.55, 0.4, [(a+"  —  ",12,SAGE_LT,True,MONO),(b,11.5,CREAM,False,BODY)])
    # code
    code_panel(s, 7.0, 3.75, 5.67, [
        [("if", C_KEY), (" (!ready || !timeUp) ", C_DEF), ("return", C_KEY), (";", C_DEF)],
        [("body.classList.", C_DEF), ("add", C_FN), ("(", C_DEF), ('"is-ready"', C_STR), (");", C_DEF)],
        [("loader.classList.", C_DEF), ("add", C_FN), ("(", C_DEF), ('"is-hidden"', C_STR), (");", C_DEF)],
        [("setTimeout", C_FN), ("(() => loader.", C_DEF), ("remove", C_FN), ("(), ", C_DEF), ("800", C_NUM), (");", C_DEF)],
    ], size=11.5, title="index.html — the reveal()")
    text(s, 7.0, 5.75, 5.67, 0.9,
         "CSS paints and animates; JavaScript only decides when the splash may leave. The two movements "
         "overlap, so Synora feels like it opens rather than loads.",
         size=11.5, color=MUTED_D, spacing=1.24)
    footer(s, 17, dark=True)
    s.notes_slide.notes_text_frame.text = (
        "The splash is a real UX + technical feature. Mechanism, honestly: the animation is pure CSS "
        "@keyframes; the exit is a CSS transition on .is-hidden; JavaScript only controls TIMING — it shows "
        "the splash for at least 2000ms, waits for DOMContentLoaded, toggles is-ready/is-hidden, then removes "
        "the node. prefers-reduced-motion is respected. Viva: is the animation JS or CSS? The motion is CSS; "
        "JS just decides when it ends.")

# =====================================================================
# 18 · BEYOND THE SYLLABUS
# =====================================================================
def s18():
    s = d.slide(BG)
    head(s, "The product we shipped", "Beyond the syllabus")
    text(s, ML, 1.62, CW, 0.5,
         "The syllabus gave us the parts — we built complete, connected features, plus personalization & "
         "avatars, light / dark themes, per-user workspaces and on-device persistence.",
         size=13.5, color=MUTED, spacing=1.2)
    cards=[("Task management","tasks.png"),("Smart calendar","calendar.png"),
           ("Attendance calculator","attendance.png"),("CGPA calculator","cgpa.png"),
           ("Achievements","achievements.png"),("Notifications","notifications.png")]
    cw=2.86; gapx=0.34; gapy=0.34
    ih=cw/1.6
    total_w=3*cw+2*gapx
    x0=(d.W-total_w)/2; y0=2.3
    rowstep=ih+0.34+gapy
    for i,(title,img) in enumerate(cards):
        col=i%3; row=i//3
        x=x0+col*(cw+gapx); y=y0+row*rowstep
        pic_card(s, IMG+img, x, y, w=cw)
        text(s, x, y+ih+0.05, cw, 0.3, title, size=11.5, color=TEXT, bold=True, align=AL.CENTER)
    footer(s, 18)
    s.notes_slide.notes_text_frame.text = (
        "This is the pay-off. Beyond the listed concepts we built a full dashboard, task manager, calendar, "
        "attendance calculator (tells you how many classes you can miss), a live CGPA calculator, an "
        "achievements system, notifications, personalization with avatars, light/dark themes, per-user "
        "workspaces and on-device persistence. Subjects are the single source of truth that attendance, CGPA "
        "and the timetable all reference.")

# =====================================================================
# 19 · WHAT WE APPLIED (SYLLABUS MAPPING)
# =====================================================================
def s19():
    s = d.slide(BG)
    head(s, "Syllabus mapping", "What we applied")
    cols=[
        ("HTML5", [("Semantic structure",1),("Forms & inputs",1),("Accessibility basics",1),("Figure / media",1)]),
        ("CSS3", [("Box model",1),("Flexbox",1),("Grid",1),("Media queries",1)]),
        ("JavaScript", [("Variables & functions",1),("Arrays & objects",1),("Loops",1),("localStorage",1)]),
        ("DOM", [("querySelector(All)",1),("getElementById",1),("create / append",1),("remove nodes",1)]),
        ("Events", [("addEventListener",1),("Delegation",1),("Form validation",1),("Capturing",0)]),
    ]
    n=len(cols); cwd=2.16; gap=(CW-n*cwd)/(n-1)
    y=2.0; ch=4.4
    for i,(hd,items) in enumerate(cols):
        x=ML+i*(cwd+gap)
        rect(s, x, y, cwd, ch, fill=SURFACE, line=MIST, lw=1.2, radius=0.07, shadow_=True, shalpha=9)
        rect(s, x, y, cwd, 0.62, fill=INK, radius=0.07)
        rect(s, x, y+0.34, cwd, 0.3, fill=INK)  # square off bottom of header
        text(s, x, y, cwd, 0.62, hd, size=14, color=CREAM, bold=True, align=AL.CENTER, anchor=AN.MIDDLE, font=HEAD)
        for j,(lab,ok) in enumerate(items):
            yy=y+0.86+j*0.82
            mark = "✓" if ok else "○"
            mc = SAGE if ok else SLATE
            dot=rect(s, x+0.22, yy, 0.34,0.34, fill=(SAGE_SOFT if ok else MIST), radius=0.4)
            text(s, x+0.22, yy-0.02, 0.34,0.34, mark, size=13, color=mc, bold=True, align=AL.CENTER, anchor=AN.MIDDLE)
            text(s, x+0.64, yy-0.06, cwd-0.78, 0.7, lab, size=11, color=(TEXT if ok else SLATE),
                 bold=False, anchor=AN.MIDDLE, spacing=1.05)
    # legend
    rich(s, ML, 6.62, CW, 0.3, [("✓  ", 12, SAGE, True, BODY),("implemented in Synora        ", 11.5, MUTED, False, BODY),
                                ("○  ", 12, SLATE, True, BODY),("learned as a concept, not a primary pattern", 11.5, MUTED, False, BODY)])
    footer(s, 19)
    s.notes_slide.notes_text_frame.text = (
        "Honest scorecard. Almost everything is ticked and genuinely implemented. The one open circle is "
        "explicit event capturing — we understand it but don't use it as a pattern; delegation via bubbling "
        "is what the app relies on. Keeping that circle open is what makes every tick credible.")

# =====================================================================
# 20 · FINAL TAKEAWAY
# =====================================================================
def s20():
    s = d.slide(INK)
    eyebrow(s, ML, 0.75, "The takeaway", color=TAUPE_LT)
    text(s, ML, 1.1, CW, 0.8, "Concepts in, product out", size=32, color=CREAM, bold=True, font=HEAD)
    # flow: syllabus -> layers -> synora
    rect(s, ML, 2.35, 3.0, 3.4, fill=INK_PANEL, radius=0.07, line=INK_SOFT, lw=1.0)
    text(s, ML, 2.6, 3.0, 0.4, "SYLLABUS", size=13, color=TAUPE_LT, bold=True, align=AL.CENTER, letter=160)
    for i,l in enumerate(["HTML5","CSS3","JavaScript","DOM","Events"]):
        text(s, ML, 3.15+i*0.46, 3.0, 0.4, l, size=15, color=CREAM, bold=True, align=AL.CENTER)
    chevron_right(s, 3.95, 4.05, size=0.42, color=TAUPE)
    rect(s, 4.65, 2.35, 3.0, 3.4, fill=INK_PANEL, radius=0.07, line=INK_SOFT, lw=1.0)
    text(s, 4.65, 2.6, 3.0, 0.4, "IMPLEMENTATION", size=13, color=TAUPE_LT, bold=True, align=AL.CENTER, letter=140)
    for i,l in enumerate(["Structure","Design system","Logic & state","Dynamic UI","Interaction"]):
        text(s, 4.65, 3.15+i*0.46, 3.0, 0.4, l, size=14, color=CREAM, align=AL.CENTER)
    chevron_right(s, 7.95, 4.05, size=0.42, color=TAUPE)
    rect(s, 8.65, 2.35, 4.0, 3.4, fill=SAGE, radius=0.07, shadow_=True, shcolor="000000", shalpha=30)
    picw(s, LOGO+"logo-mark-dark.png", 9.05, 2.72, 0.66)
    text(s, 9.8, 2.72, 3.0, 0.6, "Synora", size=24, color="FFFFFF", bold=True, font=HEAD, anchor=AN.MIDDLE)
    text(s, 9.0, 3.65, 3.4, 1.9,
         "A real, responsive student-productivity product — built entirely from the syllabus, "
         "then taken further.",
         size=14.5, color="EFF3EE", spacing=1.3)
    # closing line
    text(s, ML, 6.25, CW, 0.6, "Synora — one calm home for everything you're studying.",
         size=17, color=CREAM, bold=True, font=HEAD, align=AL.CENTER, italic=False)
    s.notes_slide.notes_text_frame.text = (
        "Close by restating the thesis as a pipeline: syllabus concepts → implementation techniques → the "
        "Synora product. We used core web-development concepts to build a real, responsive, interactive "
        "student-productivity app, then extended them. End on the tagline.")

for fn in [s1,s2,s3,s4,s5,s6,s7,s8,s9,s10,s11,s12,s13,s14,s15,s16,s17,s18,s19,s20]:
    fn()

d.save(OUT)
print("saved", OUT, "slides:", len(d.prs.slides._sldIdLst))
